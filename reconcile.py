"""Сверка заказов из CRM с банковской выпиской: что оплачено, чего не хватает и где расхождения.

Запуск:  python reconcile.py                          — демо-файлы из input/
         python reconcile.py заказы.xlsx выписка.xlsx — свои файлы
"""

import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
DATE_WINDOW = 3  # платёж без номера заказа ищем в пределах N дней от даты заказа

OK = "Совпало"
OK_NO_ID = "Совпало (без номера в назначении)"
NO_PAYMENT = "Нет оплаты в выписке"
WRONG_AMOUNT = "Сумма не совпадает"
DUPLICATE = "Двойная оплата"
NO_ORDER = "Платёж без заказа"
UNKNOWN = "Не удалось сопоставить"

FILLS = {
    OK: "C6EFCE", OK_NO_ID: "E2EFDA", NO_PAYMENT: "FFC7CE", WRONG_AMOUNT: "FFEB9C",
    DUPLICATE: "F8CBAD", NO_ORDER: "DDEBF7", UNKNOWN: "E4DFEC",
}
PROBLEMS = [NO_PAYMENT, WRONG_AMOUNT, DUPLICATE, NO_ORDER, UNKNOWN]


def load(orders_path: Path, bank_path: Path):
    orders = pd.read_excel(orders_path)
    bank = pd.read_excel(bank_path)
    orders["Дата заказа"] = pd.to_datetime(orders["Дата заказа"]).dt.date
    bank["Дата"] = pd.to_datetime(bank["Дата"]).dt.date
    # номер заказа из свободного текста: «Оплата по заказу №10452», «заказ 10452» и т.п.
    bank["Номер заказа"] = bank["Назначение платежа"].astype(str).str.extract(r"(?:№|заказ\w*)\s*(\d{4,})")[0]
    bank["Номер заказа"] = pd.to_numeric(bank["Номер заказа"], errors="coerce")
    orders["Статус сверки"], orders["Получено, ₽"], orders["Комментарий"] = "", 0.0, ""
    bank["Статус сверки"] = ""
    return orders, bank.reset_index(drop=True)


def reconcile(orders: pd.DataFrame, bank: pd.DataFrame):
    # 1. Сопоставление по номеру заказа из назначения платежа
    for oi, order in orders.iterrows():
        payments = bank[(bank["Номер заказа"] == order["Номер заказа"]) & (bank["Статус сверки"] == "")]
        if payments.empty:
            continue
        first = payments.index[0]
        received = bank.at[first, "Сумма, ₽"]
        orders.at[oi, "Получено, ₽"] = received
        if received == order["Сумма, ₽"]:
            orders.at[oi, "Статус сверки"] = bank.at[first, "Статус сверки"] = OK
        else:
            diff = received - order["Сумма, ₽"]
            orders.at[oi, "Статус сверки"] = bank.at[first, "Статус сверки"] = WRONG_AMOUNT
            orders.at[oi, "Комментарий"] = f"{'Переплата' if diff > 0 else 'Недоплата'} {abs(diff):,.0f} ₽".replace(",", " ")
        for extra in payments.index[1:]:
            bank.at[extra, "Статус сверки"] = DUPLICATE
            orders.at[oi, "Комментарий"] = f"Оплачен {len(payments)} раза"

    # 2. Платежи без номера: тот же плательщик, та же сумма, дата рядом
    for bi, pay in bank[(bank["Статус сверки"] == "") & bank["Номер заказа"].isna()].iterrows():
        candidates = orders[(orders["Статус сверки"] == "") & (orders["Клиент"] == pay["Плательщик"])
                            & (orders["Сумма, ₽"] == pay["Сумма, ₽"])]
        candidates = candidates[candidates["Дата заказа"].apply(lambda d: abs((pay["Дата"] - d).days) <= DATE_WINDOW)]
        if not candidates.empty:
            oi = candidates.index[0]
            orders.at[oi, "Статус сверки"] = bank.at[bi, "Статус сверки"] = OK_NO_ID
            orders.at[oi, "Получено, ₽"] = pay["Сумма, ₽"]
            orders.at[oi, "Комментарий"] = f"Найден по сумме и плательщику: «{pay['Назначение платежа']}»"

    # 3. Всё, что осталось без пары
    orders.loc[orders["Статус сверки"] == "", "Статус сверки"] = NO_PAYMENT
    leftover = bank["Статус сверки"] == ""
    bank.loc[leftover & bank["Номер заказа"].notna(), "Статус сверки"] = NO_ORDER
    bank.loc[leftover & bank["Номер заказа"].isna(), "Статус сверки"] = UNKNOWN
    return orders, bank


def problems_table(orders, bank) -> pd.DataFrame:
    rows = []
    for _, o in orders[orders["Статус сверки"].isin(PROBLEMS)].iterrows():
        rows.append({"Проблема": o["Статус сверки"], "Дата": o["Дата заказа"], "Клиент": o["Клиент"],
                     "Номер заказа": o["Номер заказа"], "Ожидалось, ₽": o["Сумма, ₽"],
                     "Получено, ₽": o["Получено, ₽"], "Комментарий": o["Комментарий"]})
    for _, b in bank[bank["Статус сверки"].isin([DUPLICATE, NO_ORDER, UNKNOWN])].iterrows():
        rows.append({"Проблема": b["Статус сверки"], "Дата": b["Дата"], "Клиент": b["Плательщик"],
                     "Номер заказа": "" if pd.isna(b["Номер заказа"]) else int(b["Номер заказа"]),
                     "Ожидалось, ₽": "", "Получено, ₽": b["Сумма, ₽"],
                     "Комментарий": f"Назначение: «{b['Назначение платежа']}»"})
    df = pd.DataFrame(rows)
    return df.sort_values(["Проблема", "Дата"], key=lambda s: s.map(PROBLEMS.index) if s.name == "Проблема" else s) if rows else df


# ---------- Excel ----------

def _sheet(wb, title, df, widths, status_col=None, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font = PatternFill("solid", fgColor="1F4E78"), Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    cols = list(df.columns)
    for r in range(2, ws.max_row + 1):
        if status_col:
            cell = ws.cell(row=r, column=cols.index(status_col) + 1)
            cell.fill = PatternFill("solid", fgColor=FILLS.get(cell.value, "FFFFFF"))
        for m in money_cols:
            ws.cell(row=r, column=cols.index(m) + 1).number_format = '# ##0 "₽"'
        for c in ws[r]:
            if hasattr(c.value, "year"):
                c.number_format = "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _summary(wb, orders, bank):
    ws = wb.create_sheet("Итоги", 0)
    title = Font(bold=True, size=14)
    ws["A1"], ws["A1"].font = "Сверка заказов с банковской выпиской", title
    expected, received = orders["Сумма, ₽"].sum(), bank["Сумма, ₽"].sum()
    lines = [
        ("Заказов в CRM", len(orders)), ("Платежей в выписке", len(bank)), ("", ""),
        ("Сумма заказов, ₽", expected), ("Поступило по выписке, ₽", received),
        ("Разница, ₽", received - expected), ("", ""), ("Результат по заказам", ""),
    ]
    for status in [OK, OK_NO_ID, WRONG_AMOUNT, NO_PAYMENT]:
        lines.append((status, int((orders["Статус сверки"] == status).sum())))
    lines += [("", ""), ("Платежи без пары", "")]
    for status in [DUPLICATE, NO_ORDER, UNKNOWN]:
        lines.append((status, int((bank["Статус сверки"] == status).sum())))
    for i, (k, v) in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=k)
        cell = ws.cell(row=i, column=2, value=v)
        if k in FILLS:
            ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=FILLS[k])
        if "₽" in k:
            cell.number_format = '# ##0 "₽"'
        if k in ("Результат по заказам", "Платежи без пары"):
            ws.cell(row=i, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 38, 16


def save_report(orders, bank, problems, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _summary(wb, orders, bank)
    _sheet(wb, "Расхождения", problems, [26, 12, 22, 14, 14, 14, 60], "Проблема", ["Ожидалось, ₽", "Получено, ₽"])
    _sheet(wb, "Заказы", orders, [14, 12, 22, 12, 11, 30, 13, 50], "Статус сверки", ["Сумма, ₽", "Получено, ₽"])
    bank_view = bank[["Дата", "Плательщик", "Сумма, ₽", "Назначение платежа", "Номер заказа", "Статус сверки"]]
    _sheet(wb, "Выписка", bank_view, [12, 22, 12, 32, 14, 36], "Статус сверки", ["Сумма, ₽"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main():
    orders_path = Path(sys.argv[1]) if len(sys.argv) > 2 else ROOT / "input" / "orders.xlsx"
    bank_path = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "input" / "bank.xlsx"
    orders, bank = reconcile(*load(orders_path, bank_path))
    problems = problems_table(orders, bank)
    out = save_report(orders, bank, problems, ROOT / "output" / f"sverka_{datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx")
    print(f"Заказов: {len(orders)}, платежей: {len(bank)}, расхождений: {len(problems)}")
    print(problems["Проблема"].value_counts().to_string() if len(problems) else "Расхождений нет")
    print(f"Отчёт: {out}")


if __name__ == "__main__":
    main()
